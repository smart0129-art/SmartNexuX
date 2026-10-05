import io
import tempfile
import unittest
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SpreadsheetImage
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.services.document_images import (
    associate_images_with_chunks,
    extract_document_images,
    split_pptx_slides,
    split_xlsx_sheets,
)
from app.services.document_parser import DocumentChunk


def sample_png() -> bytes:
    stream = io.BytesIO()
    Image.new("RGB", (32, 24), color="navy").save(stream, format="PNG")
    return stream.getvalue()


class DocumentImageTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary_directory.name)

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def test_pdf_page_images_are_associated_with_page_chunks(self) -> None:
        source = self.root / "report.pdf"
        document = pymupdf.open()
        page = document.new_page()
        page.insert_image(
            pymupdf.Rect(20, 20, 120, 100),
            stream=sample_png(),
        )
        document.save(source)
        document.close()

        images = extract_document_images(source)
        chunks = [
            DocumentChunk(
                chunk_index=0,
                content="Chart data",
                metadata={"page_number": 1},
            ),
            DocumentChunk(
                chunk_index=1,
                content="Other content",
                metadata={"page_number": 2},
            ),
        ]
        associate_images_with_chunks(chunks, images, ".pdf")

        self.assertEqual(len(images), 1)
        self.assertEqual(
            chunks[0].metadata["image_asset_ids"],
            [images[0].image_id],
        )
        self.assertNotIn("image_asset_ids", chunks[1].metadata)

    def test_docx_image_is_associated_only_with_its_ocr_chunk(self) -> None:
        source = self.root / "report.docx"
        document = Document()
        document.add_picture(io.BytesIO(sample_png()))
        document.save(source)

        images = extract_document_images(source)
        chunk = DocumentChunk(
            chunk_index=0,
            content="*[Image OCR]\nChart title\n[End OCR]*",
            metadata={},
        )
        associate_images_with_chunks([chunk], images, ".docx")

        self.assertEqual(len(images), 1)
        self.assertEqual(
            chunk.metadata["image_asset_ids"],
            [images[0].image_id],
        )

    def test_pptx_image_is_associated_with_its_slide(self) -> None:
        source = self.root / "slides.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(
            io.BytesIO(sample_png()),
            Inches(1),
            Inches(1),
        )
        presentation.save(source)

        images = extract_document_images(source)
        chunk = DocumentChunk(
            chunk_index=0,
            content="Revenue by quarter",
            metadata={"slide_number": 1},
        )
        associate_images_with_chunks([chunk], images, ".pptx")

        self.assertEqual(len(images), 1)
        self.assertEqual(
            chunk.metadata["image_asset_ids"],
            [images[0].image_id],
        )

    def test_xlsx_images_are_associated_with_the_matching_sheet_ocr(self) -> None:
        source = self.root / "workbook.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Revenue"
        sheet.add_image(SpreadsheetImage(io.BytesIO(sample_png())), "A1")
        workbook.save(source)

        images = extract_document_images(source)
        chunk = DocumentChunk(
            chunk_index=0,
            content="*[Image OCR]\nQuarterly chart\n[End OCR]*",
            metadata={"sheet_name": "Revenue"},
        )
        associate_images_with_chunks([chunk], images, ".xlsx")

        self.assertEqual(len(images), 1)
        self.assertEqual(
            chunk.metadata["image_asset_ids"],
            [images[0].image_id],
        )

    def test_slide_and_sheet_markdown_are_split_with_source_locations(self) -> None:
        slides = split_pptx_slides(
            "<!-- Slide number: 1 -->\nFirst slide\n"
            "<!-- Slide number: 2 -->\nSecond slide"
        )
        sheets = split_xlsx_sheets(
            "## Revenue\nQuarterly totals\n\n## Costs\nBudget totals"
        )

        self.assertEqual(slides, [(1, "First slide"), (2, "Second slide")])
        self.assertEqual(
            sheets,
            [("Revenue", "Quarterly totals"), ("Costs", "Budget totals")],
        )


if __name__ == "__main__":
    unittest.main()
